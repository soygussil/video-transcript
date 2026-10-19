import json
from pathlib import Path

import cv2
import numpy as np
import pytest
from docx import Document
from openpyxl import load_workbook

from analizador_video.application.export_reports import apply_corrections
from analizador_video.config.settings import load_settings
from analizador_video.domain.alignment import align
from analizador_video.domain.enums import ReportKind, TranscriptSource
from analizador_video.domain.models import ScreenCapture, TranscriptSegment, UserCorrection, VideoMetadata
from analizador_video.infrastructure.logging_setup import sanitize
from analizador_video.infrastructure.reports.model import EXCEL_CELL_LIMIT, ReportData
from analizador_video.infrastructure.reports.service import ReportService
from analizador_video.infrastructure.video.frames import KeyframeDetector

VIDEO = load_settings().video


def frame(color):
    return np.full((240, 320, 3), color, np.uint8)


def test_detector_deduplicates_identical_frames_and_detects_changes():
    detector = KeyframeDetector(VIDEO)
    accepted = [
        (r.time, r.reason)
        for s in range(12)
        if (r := detector.feed(float(s), frame((200, 60, 40) if s < 6 else (40, 60, 200))))
    ]
    # Comportamiento original conservado: se guarda el primer cuadro estable tras el cambio.
    assert accepted == [(0.0, "inicial"), (7.0, "cambio_estable")]


def test_detector_respects_minimum_separation():
    detector = KeyframeDetector(VIDEO)
    colors = [(200, 0, 0), (0, 200, 0), (0, 0, 200), (200, 200, 0)]
    assert len([r for t, c in enumerate(colors) if (r := detector.feed(t * 0.5, frame(c)))]) == 1


def _data(tmp_path, text="Hola", segments=None, sensitive=True):
    image = tmp_path / "frames" / "P0001.jpg"
    image.parent.mkdir(parents=True, exist_ok=True)
    cv2.imwrite(str(image), frame((10, 120, 200)))
    capture = ScreenCapture("P0001", 1, "frames/P0001.jpg", 0.0, 1.0, "inicial")
    segs = (
        segments if segments is not None else [TranscriptSegment(1, 0.5, 1.5, "Ana", text, TranscriptSource.IMPORTED)]
    )
    return ReportData(
        "Demo",
        "demo.mp4",
        VideoMetadata(10.0, 25.0, 250, 320, 240, "mp4v"),
        "Balanceado",
        "Importada",
        align([capture], segs, {"P0001": "OCR"}, 10.0),
        segs,
        tmp_path,
        "2026-10-19 10:00",
        "0.2.0",
        "Reporte",
        "Confidencial",
        warnings=["Advertencia de prueba"],
        include_sensitive=sensitive,
    )


def test_reports_minimal_data(tmp_path):
    files = ReportService().generate(_data(tmp_path), list(ReportKind), tmp_path / "reports")
    assert sorted(a.path.name for a in files) == [
        "analisis.json",
        "reporte_analisis.docx",
        "reporte_analisis.xlsx",
        "transcripcion.srt",
        "transcripcion.txt",
    ]  # problema G
    wb = load_workbook(tmp_path / "reports" / "reporte_analisis.xlsx")
    assert wb.sheetnames == [
        "Resumen",
        "Pantallas",
        "Visto vs. dicho",
        "Segundo a segundo",
        "Transcripción",
        "Advertencias",
    ]
    assert wb["Pantallas"].freeze_panes == "A2" and wb["Pantallas"]["B2"].number_format == "[h]:mm:ss.000"
    text = "\n".join(p.text for p in Document(tmp_path / "reports" / "reporte_analisis.docx").paragraphs)
    assert "Resumen de ejecución" in text and "Transcripción completa" in text
    data = json.loads((tmp_path / "reports" / "analisis.json").read_text(encoding="utf-8"))
    assert data["schema_version"] == 2 and data["capturas"][0]["segmentos"] == [1]
    assert data["segmentos"][0]["start_seconds"] == 0.5


def test_reports_without_transcript(tmp_path):
    assert len(ReportService().generate(_data(tmp_path, segments=[]), list(ReportKind), tmp_path / "r")) == 5
    assert (tmp_path / "r" / "transcripcion.srt").read_text(encoding="utf-8") == ""


def test_reports_long_text_and_control_chars(tmp_path):
    ReportService().generate(
        _data(tmp_path, text=("palabra \x07" * 6000).strip()), [ReportKind.WORD, ReportKind.EXCEL], tmp_path / "r"
    )
    value = load_workbook(tmp_path / "r" / "reporte_analisis.xlsx")["Transcripción"]["E2"].value
    assert len(value) <= EXCEL_CELL_LIMIT + 30 and "\x07" not in value


def test_reports_exclude_sensitive(tmp_path):
    ReportService().generate(_data(tmp_path, sensitive=False), [ReportKind.SRT], tmp_path / "r")
    srt = (tmp_path / "r" / "transcripcion.srt").read_text(encoding="utf-8")
    assert "Ana" not in srt and "Participante" in srt


def test_corrections_latest_wins():
    segments = [TranscriptSegment(1, 0, 1, "A", "hola", TranscriptSource.IMPORTED)]
    corrections = [
        UserCorrection("segmento", "1", "text", "hola", "Hola", "t1"),
        UserCorrection("segmento", "1", "text", "Hola", "Hola, equipo", "t2"),
        UserCorrection("segmento", "1", "speaker", "A", "Ana", "t3"),
        UserCorrection("ocr", "P0001", "text", "Guadar", "Guardar", "t4"),
    ]
    fixed, ocr = apply_corrections(segments, {"P0001": "Guadar"}, corrections)
    assert fixed[0].text == "Hola, equipo" and fixed[0].speaker == "Ana" and ocr["P0001"] == "Guardar"


@pytest.mark.parametrize(
    "raw,leak",
    [
        ("api_key=sk-123456 abc", "sk-123456"),
        ("Authorization: Bearer eyJabc.def", "eyJabc"),
        ("token = secreto", "secreto"),
    ],
)
def test_sanitize_secrets(raw, leak):
    assert leak not in sanitize(raw)


def test_sanitize_home_path():
    assert str(Path.home()) not in sanitize(f"Error en {Path.home()}/Videos/demo.mp4")
