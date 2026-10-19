import os

import pytest

from analizador_video.application.analyze_video import AnalyzeVideo
from analizador_video.application.backends import BackendRegistry
from analizador_video.application.events import JobFailed, JobStarted, PartialResult, StageFinished, WarningRaised
from analizador_video.application.export_reports import ExportReports
from analizador_video.application.job_manager import new_run_id
from analizador_video.application.project_service import CreateProjectRequest, ProjectService
from analizador_video.domain.enums import DeviceKind, JobStatus, ProfileName, ReportKind, SourceMode, StageName
from analizador_video.domain.errors import ProjectBusyError
from analizador_video.domain.models import StagePlan, UserCorrection
from analizador_video.infrastructure.hardware.compatibility import NOT_COMPATIBLE, CompatibilityStore
from analizador_video.infrastructure.persistence.repository import ProjectRepository
from tests.fixtures.fakes import CountingToken, FakeOcrBackend, FakeTranscriptionBackend
from tests.fixtures.media import add_silent_audio

ALL = frozenset(StageName.ordered())
FRAMES_ONLY = StagePlan(frozenset({StageName.PREPARE, StageName.FRAMES}))


def create(projects_root, video, transcript=None, mode=SourceMode.REFERENCE, project_id=None):
    return ProjectService(projects_root).create(
        CreateProjectRequest(video, transcript, ProfileName.BALANCED, mode, project_id=project_id)
    )


def run(root, plan=None, registry=None, token=None):
    events = []
    status = AnalyzeVideo(registry or BackendRegistry()).run(
        root, plan or StagePlan(ALL), events.append, token or CountingToken(), new_run_id()
    )
    return status, events


def set_config(root, section, key, value):
    repo = ProjectRepository.open(root)
    config = repo.get_config()
    config[section][key] = value
    repo.set_config(config)
    repo.close()


def test_full_run_with_vtt_and_fake_ocr(projects_root, sample_video, sample_vtt):
    project = create(projects_root, sample_video, sample_vtt)
    status, events = run(project.root, registry=BackendRegistry(ocr={"fake_ocr": FakeOcrBackend()}))
    assert status is JobStatus.COMPLETED
    repo = ProjectRepository.open(project.root)
    assert [round(c.time_seconds) for c in repo.captures()] == [0, 5, 9]
    segments = repo.segments()
    assert len(segments) == 3 and segments[0].speaker == "Ana López"
    assert repo.ocr_results()["P0001"].text == "Botón Guardar"  # baja confianza filtrada
    assert {r.kind for r in repo.reports()} == set(ReportKind)
    run_id = next(e.run_id for e in events if isinstance(e, JobStarted))
    assert {m["stage"] for m in repo.metrics(run_id)} >= {"pantallas", "ocr"}
    repo.close()
    assert sum(isinstance(e, PartialResult) for e in events) == 3
    for name in ("reporte_analisis.docx", "reporte_analisis.xlsx", "analisis.json", "transcripcion.srt"):
        assert (project.root / "reports" / name).exists()


def test_reports_have_second_by_second_timeline(projects_root, sample_video, sample_vtt):
    import docx
    import openpyxl

    project = create(projects_root, sample_video, sample_vtt)
    run(project.root, registry=BackendRegistry(ocr={"fake_ocr": FakeOcrBackend()}))
    sheet = openpyxl.load_workbook(project.root / "reports" / "reporte_analisis.xlsx")["Segundo a segundo"]
    rows = list(sheet.iter_rows(min_row=2, values_only=True))  # segundo, tiempo, dice, pantalla, OCR
    assert [r[0] for r in rows] == list(range(12))
    assert "Bienvenidos" in rows[0][2] and "Bienvenidos" in rows[2][2] and not rows[3][2]  # 0.5–3.0 s
    assert "segunda pantalla" in rows[4][2] and "segunda pantalla" in rows[5][2]
    assert rows[0][3] == "P0001" and rows[5][3] == "P0002" and rows[9][3] == "P0003"
    assert rows[0][4] == "Botón Guardar" == rows[3][4]  # el OCR de la pantalla vigente acompaña cada segundo
    word = docx.Document(project.root / "reports" / "reporte_analisis.docx")
    timeline = next(t for t in word.tables if t.rows[0].cells[0].text == "Tiempo")
    assert len(timeline.rows) == 13 and timeline.rows[1].cells[3].text == "Botón Guardar"
    assert timeline.rows[2].cells[3].text == ""  # el OCR no se repite mientras la pantalla no cambia


def test_without_ocr_and_without_automatic_transcription(projects_root, sample_video):
    project = create(projects_root, sample_video)
    status, events = run(project.root, StagePlan(ALL, automatic_transcription=False, run_ocr=False))
    finished = {e.stage: e.status for e in events if isinstance(e, StageFinished)}
    assert finished["ocr"] == "skipped" and finished["audio"] == "skipped" and status is JobStatus.COMPLETED


def test_missing_ocr_backend_is_a_warning_not_a_failure(projects_root, sample_video):
    project = create(projects_root, sample_video)
    status, events = run(
        project.root,
        StagePlan(ALL, automatic_transcription=False),
        BackendRegistry(ocr={"f": FakeOcrBackend(installed=False)}),
    )
    assert status is JobStatus.COMPLETED_WITH_WARNINGS
    assert any(isinstance(e, WarningRaised) and e.stage == "ocr" for e in events)


def test_cancel_and_resume_frames_gives_same_captures(projects_root, video_factory):
    video = video_factory("largo.mp4", scenes=4, seconds_per_scene=5)
    reference = create(projects_root, video)
    run(reference.root, FRAMES_ONLY)
    expected = [(c.id, c.time_seconds) for c in ProjectRepository.open(reference.root).captures()]
    project = create(projects_root, video, project_id="reanudable")
    assert run(project.root, FRAMES_ONLY, token=CountingToken(after=9))[0] is JobStatus.CANCELLED
    repo = ProjectRepository.open(project.root)
    partial = repo.get_stage(StageName.FRAMES)
    assert partial.status.value == "cancelled" and partial.checkpoint["next_index"] > 0
    assert 0 < len(repo.captures()) < len(expected)
    repo.close()
    assert run(project.root, FRAMES_ONLY)[0] is JobStatus.COMPLETED
    assert [(c.id, c.time_seconds) for c in ProjectRepository.open(project.root).captures()] == expected


def test_ocr_checkpoints_and_resume(projects_root, video_factory):  # problemas D y E
    project = create(projects_root, video_factory("muchas.mp4", scenes=4, seconds_per_scene=3))
    run(project.root, FRAMES_ONLY)
    plan = StagePlan(frozenset({StageName.PREPARE, StageName.FRAMES, StageName.OCR}))
    status, _ = run(project.root, plan, BackendRegistry(ocr={"f": FakeOcrBackend()}), CountingToken(after=3))
    assert status is JobStatus.CANCELLED
    repo = ProjectRepository.open(project.root)
    saved, total = len(repo.ocr_results()), len(repo.captures())
    repo.close()
    assert 0 < saved < total
    second = FakeOcrBackend()
    assert run(project.root, plan, BackendRegistry(ocr={"f": second}))[0] is JobStatus.COMPLETED
    assert second.reader.calls == total - saved  # sólo procesa lo pendiente


def test_ocr_systemic_failure(projects_root, sample_video):
    project = create(projects_root, sample_video)
    set_config(project.root, "ocr", "max_errores_consecutivos", 2)
    backend = FakeOcrBackend()
    backend.reader.fail_on = set(range(1, 100))
    status, events = run(project.root, registry=BackendRegistry(ocr={"f": backend}))
    failure = next(e for e in events if isinstance(e, JobFailed))
    assert status is JobStatus.FAILED and failure.message["code"] == "ocr_sistemico" and failure.message["saved"]


def test_report_requires_alignment(projects_root, sample_video):
    project = create(projects_root, sample_video)
    status, events = run(project.root, StagePlan(frozenset({StageName.PREPARE, StageName.REPORT})))
    assert status is JobStatus.FAILED
    assert next(e for e in events if isinstance(e, JobFailed)).message["code"] == "dependencia_de_etapa"


def test_new_ocr_invalidates_previous_alignment(projects_root, sample_video):  # problema C
    project = create(projects_root, sample_video)
    run(project.root, StagePlan(ALL, run_ocr=False, automatic_transcription=False))
    status, _ = run(
        project.root, StagePlan(ALL, automatic_transcription=False), BackendRegistry(ocr={"f": FakeOcrBackend()})
    )
    assert status is JobStatus.COMPLETED
    assert "Botón Guardar" in (project.root / "reports" / "analisis.json").read_text(encoding="utf-8")


def test_source_missing_and_changed(projects_root, video_factory, tmp_path):
    video = video_factory("mover.mp4")
    project = create(projects_root, video)
    run(project.root, StagePlan(frozenset({StageName.PREPARE})))
    video.rename(tmp_path / "otro.mp4")
    status, events = run(project.root)
    assert status is JobStatus.FAILED
    assert next(e for e in events if isinstance(e, JobFailed)).message["code"] == "video_no_encontrado"
    video_factory("mover.mp4", scenes=2)
    _, events = run(project.root)
    assert next(e for e in events if isinstance(e, JobFailed)).message["code"] == "entrada_invalida"


def test_copy_mode_survives_original_removal(projects_root, video_factory):
    video = video_factory("copiar.mp4")
    project = create(projects_root, video, mode=SourceMode.COPY)
    run(project.root, StagePlan(frozenset({StageName.PREPARE})))
    video.unlink()
    assert run(project.root, FRAMES_ONLY)[0] is JobStatus.COMPLETED
    assert (project.root / "source" / "copiar.mp4").exists()


def test_project_lock_prevents_parallel_runs(projects_root, sample_video):
    project = create(projects_root, sample_video)
    # El proceso padre siempre está vivo y es distinto del actual (en Windows el PID 1 no existe).
    (project.root / ".en_proceso.lock").write_text(str(os.getppid()), encoding="utf-8")
    with pytest.raises(ProjectBusyError):
        run(project.root)


def test_export_without_reprocessing_applies_corrections(projects_root, sample_video, sample_vtt):
    project = create(projects_root, sample_video, sample_vtt)
    run(project.root, StagePlan(ALL - {StageName.REPORT}, run_ocr=False))
    repo = ProjectRepository.open(project.root)
    repo.add_correction(UserCorrection("segmento", "1", "text", "x", "Texto corregido", ""))
    repo.close()
    artifacts = ExportReports().run(project.root, [ReportKind.SRT], project.root / "otra")
    assert "Texto corregido" in artifacts[0].path.read_text(encoding="utf-8")


@pytest.fixture
def video_with_audio(sample_video, tmp_path):
    return add_silent_audio(sample_video, tmp_path / "con_audio.mp4")


def test_automatic_transcription_gpu_failure_falls_back_to_cpu(projects_root, video_with_audio, tmp_path):
    ov = FakeTranscriptionBackend("openvino", (DeviceKind.GPU, DeviceKind.CPU), failing_devices={DeviceKind.GPU})
    fw = FakeTranscriptionBackend("faster_whisper", (DeviceKind.CPU,))
    store = CompatibilityStore(tmp_path / "compat.json")
    project = create(projects_root, video_with_audio)
    status, events = run(
        project.root,
        StagePlan(ALL, run_ocr=False, device=DeviceKind.GPU),
        BackendRegistry({"openvino": ov, "faster_whisper": fw}, {}, store, "2025.1"),
    )
    assert status is JobStatus.COMPLETED_WITH_WARNINGS
    assert ProjectRepository.open(project.root).segments()[0].text == "texto cpu"
    assert any("GPU" in e.message for e in events if isinstance(e, WarningRaised))
    assert store.load()[0].status == NOT_COMPATIBLE  # Automático no volverá a usarla


def test_ffmpeg_missing_is_actionable_warning(projects_root, video_with_audio, monkeypatch):
    import analizador_video.infrastructure.audio as audio

    monkeypatch.setattr(audio, "find_ffmpeg", lambda: None)
    monkeypatch.setitem(__import__("sys").modules, "av", None)  # sin FFmpeg ni PyAV
    project = create(projects_root, video_with_audio)
    status, events = run(
        project.root,
        StagePlan(ALL, run_ocr=False),
        BackendRegistry({"faster_whisper": FakeTranscriptionBackend("faster_whisper", (DeviceKind.CPU,))}),
    )
    warning = next(e for e in events if isinstance(e, WarningRaised) and e.code == "ffmpeg_no_encontrado")
    assert "FFmpeg" in warning.message and "VTT o SRT" in warning.message and "guardado" in warning.message
    assert status is JobStatus.COMPLETED_WITH_WARNINGS


def test_audio_extraction_works_without_ffmpeg_binary(video_with_audio, tmp_path, monkeypatch):
    import analizador_video.infrastructure.audio as audio

    monkeypatch.setattr(audio, "find_ffmpeg", lambda: None)
    wav = audio.extract_audio(video_with_audio, tmp_path / "a.wav", lambda: False)
    assert audio.wav_duration_seconds(wav) == pytest.approx(12, abs=0.5)
    assert abs(audio.read_wav_block(wav, 1, 2)).max() > 0.1  # el tono se conservó


def test_audio_extraction_without_ffmpeg_reports_missing_stream(sample_video, tmp_path, monkeypatch):
    import analizador_video.infrastructure.audio as audio

    monkeypatch.setattr(audio, "find_ffmpeg", lambda: None)
    with pytest.raises(audio.NoAudioStreamError):
        audio.extract_audio(sample_video, tmp_path / "a.wav", lambda: False)


def test_video_without_audio(projects_root, sample_video):
    project = create(projects_root, sample_video)
    status, events = run(
        project.root,
        StagePlan(ALL, run_ocr=False),
        BackendRegistry({"faster_whisper": FakeTranscriptionBackend("faster_whisper", (DeviceKind.CPU,))}),
    )
    assert any(isinstance(e, WarningRaised) and e.code == "sin_audio" for e in events)
    assert status is JobStatus.COMPLETED_WITH_WARNINGS


def test_transcription_blocks_use_bloque_minutos(projects_root, video_with_audio):  # problema H
    project = create(projects_root, video_with_audio)
    set_config(project.root, "transcript", "bloque_minutos", 1.0)
    run(
        project.root,
        StagePlan(ALL, run_ocr=False),
        BackendRegistry({"faster_whisper": FakeTranscriptionBackend("faster_whisper", (DeviceKind.CPU,))}),
    )
    state = ProjectRepository.open(project.root).get_stage(StageName.TRANSCRIPT)
    assert state.checkpoint == {"next_block": 1} and state.total == 1
    assert (project.root / "audio" / "audio_16k.wav").exists()
