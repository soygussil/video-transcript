"""Reporte Excel `reporte_analisis.xlsx` a partir de datos estructurados."""

from __future__ import annotations

import os
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.worksheet import Worksheet

from analizador_video.domain.timefmt import format_hms
from analizador_video.infrastructure.reports.model import (
    EXCEL_CELL_LIMIT,
    ReportData,
    build_timeline,
    safe_text,
    speaker_label,
)

EXCEL_REPORT_NAME = "reporte_analisis.xlsx"
TIME_FORMAT = "[h]:mm:ss.000"
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")


def _cell_text(value: str) -> str:
    text = safe_text(value)
    return text if len(text) <= EXCEL_CELL_LIMIT else text[:EXCEL_CELL_LIMIT] + " […texto recortado]"


def _excel_time(seconds: float) -> float:
    return round(seconds, 3) / 86400.0


def _style(ws: Worksheet, time_columns: tuple[int, ...] = (), widths: dict[int, int] | None = None) -> None:
    ws.freeze_panes = "A2"
    if ws.max_row > 1:
        ws.auto_filter.ref = ws.dimensions
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for column in ws.iter_cols(min_row=1):
        index = column[0].column
        longest = max((len(str(c.value or "")) for c in column[:200]), default=10)
        ws.column_dimensions[column[0].column_letter].width = (widths or {}).get(index, min(60, max(12, longest + 2)))
        for cell in column[1:]:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if index in time_columns:
                cell.number_format = TIME_FORMAT


def create_excel(data: ReportData, destination: Path) -> Path:
    wb = Workbook()
    summary = wb.active
    summary.title = "Resumen"
    summary.append(["Concepto", "Valor"])
    rows: list[tuple[str, object]] = [
        ("Proyecto", data.project_name),
        ("Video", data.video_name),
        ("Capturas", len(data.evidence)),
        ("Segmentos de transcripción", len(data.segments)),
        ("Origen de la transcripción", data.transcript_source_label),
        ("Perfil", data.profile_label),
        ("Correcciones aplicadas", len(data.corrections)),
        ("Generado", data.generated_at),
        ("Versión de la aplicación", data.app_version),
    ]
    if data.metadata is not None:
        rows += [
            ("Duración", format_hms(data.metadata.duration_seconds)),
            ("Resolución", f"{data.metadata.width} × {data.metadata.height}"),
        ]
    rows += list(data.settings_summary.items())
    for key, value in rows:
        summary.append([key, value])
    _style(summary, widths={1: 32, 2: 60})

    screens = wb.create_sheet("Pantallas")
    screens.append(["ID", "Inicio", "Fin", "Captura", "Cambio", "Motivo", "OCR"])
    for item in data.evidence:
        screens.append(
            [
                item.capture.id,
                _excel_time(item.start_seconds),
                _excel_time(item.end_seconds),
                Path(item.capture.file_rel).name,
                item.capture.change_score,
                item.capture.reason,
                _cell_text(item.ocr_text),
            ]
        )
        if data.include_sensitive:
            link = screens.cell(row=screens.max_row, column=4)
            link.hyperlink = Path(os.path.relpath(data.project_root / item.capture.file_rel, destination)).as_posix()
            link.font = Font(color="0563C1", underline="single")
    _style(screens, time_columns=(2, 3), widths={7: 60})

    seen = wb.create_sheet("Visto vs. dicho")
    seen.append(["ID", "Inicio", "Fin", "Texto visible (OCR)", "Lo que se dijo"])
    for item in data.evidence:
        spoken = "\n".join(
            f"[{s.start_display} - {s.end_display}] {speaker_label(s, data.include_sensitive)}: {s.text}"
            for s in item.segments
        )
        seen.append(
            [
                item.capture.id,
                _excel_time(item.start_seconds),
                _excel_time(item.end_seconds),
                _cell_text(item.ocr_text),
                _cell_text(spoken),
            ]
        )
    _style(seen, time_columns=(2, 3), widths={4: 50, 5: 70})

    timeline = wb.create_sheet("Segundo a segundo")
    timeline.append(["Segundo", "Tiempo", "Lo que se dice", "Pantalla", "Texto visible (OCR)"])
    for row in build_timeline(data):
        timeline.append(
            [row.second, _excel_time(row.second), _cell_text(row.spoken), row.capture_id, _cell_text(row.ocr_text)]
        )
    _style(timeline, time_columns=(2,), widths={1: 10, 2: 14, 3: 70, 4: 10, 5: 50})

    transcript = wb.create_sheet("Transcripción")
    transcript.append(["#", "Inicio", "Fin", "Hablante", "Texto", "Confianza"])
    for s in data.segments:
        transcript.append(
            [
                s.seq,
                _excel_time(s.start_seconds),
                _excel_time(s.end_seconds),
                speaker_label(s, data.include_sensitive),
                _cell_text(s.text),
                s.confidence,
            ]
        )
    _style(transcript, time_columns=(2, 3), widths={5: 80})

    warnings = wb.create_sheet("Advertencias")
    warnings.append(["Advertencia"])
    for warning in data.warnings:
        warnings.append([_cell_text(warning)])
    _style(warnings, widths={1: 100})

    destination.mkdir(parents=True, exist_ok=True)
    path = destination / EXCEL_REPORT_NAME
    temporary = destination / f".{EXCEL_REPORT_NAME}.tmp"
    wb.save(str(temporary))
    temporary.replace(path)
    return path
