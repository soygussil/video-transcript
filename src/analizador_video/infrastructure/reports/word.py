"""Reporte Word `reporte_analisis.docx` a partir de datos estructurados."""

from __future__ import annotations

from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, UnidentifiedImageError

from analizador_video.domain.timefmt import format_hms
from analizador_video.infrastructure.reports.model import ReportData, build_timeline, safe_text, speaker_label

WORD_REPORT_NAME = "reporte_analisis.docx"
MAX_W_IN, MAX_H_IN = 7.0, 5.2
ACCENT = RGBColor(0x1F, 0x4E, 0x79)


def _image_size(path: Path) -> tuple[float, float] | None:
    try:
        with Image.open(path) as image:
            width, height = image.size
    except (OSError, UnidentifiedImageError):
        return None
    scale = min(MAX_W_IN / width, MAX_H_IN / height)  # sin deformar
    return width * scale, height * scale


def _add_toc(paragraph: Any) -> None:
    run = paragraph.add_run()
    parts = []
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = 'TOC \\o "1-2" \\h \\z \\u'
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    placeholder = OxmlElement("w:t")
    placeholder.text = "Haz clic derecho y elige «Actualizar campo» para generar el índice."
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    parts = [begin, instr, separate, placeholder, end]
    for element in parts:
        run._r.append(element)


def _key_value_table(doc: Any, rows: list[tuple[str, str]]) -> None:
    table = doc.add_table(rows=0, cols=2)
    table.style = "Light List Accent 1"
    for key, value in rows:
        cells = table.add_row().cells
        cells[0].text = key
        cells[1].text = safe_text(value)
        cells[0].paragraphs[0].runs[0].bold = True


def _segment_paragraph(paragraph: Any, segment: Any, include_sensitive: bool, color: bool = False) -> None:
    paragraph.add_run(f"[{segment.start_display} – {segment.end_display}] ").bold = True
    paragraph.add_run(f"{safe_text(speaker_label(segment, include_sensitive))}: ").bold = True
    text = paragraph.add_run(safe_text(segment.text))
    if color:
        text.font.color.rgb = ACCENT


def create_word(data: ReportData, destination: Path) -> Path:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = section.bottom_margin = Inches(0.6)
    section.left_margin = section.right_margin = Inches(0.6)
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(10.5)
    section.header.paragraphs[0].text = safe_text(data.header or data.title)
    footer = section.footer.paragraphs[0]
    footer.text = safe_text(" · ".join(x for x in (data.footer, data.confidentiality) if x))
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER

    for _ in range(6):
        doc.add_paragraph()
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run(safe_text(data.title))
    run.bold = True
    run.font.size = Pt(24)
    run.font.color.rgb = ACCENT
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run(safe_text(data.project_name)).font.size = Pt(14)
    if data.confidentiality:
        conf = doc.add_paragraph()
        conf.alignment = WD_ALIGN_PARAGRAPH.CENTER
        conf.add_run(safe_text(data.confidentiality)).bold = True
    date = doc.add_paragraph()
    date.alignment = WD_ALIGN_PARAGRAPH.CENTER
    date.add_run(f"Generado: {data.generated_at}")
    doc.add_page_break()

    doc.add_heading("Contenido", level=1)
    _add_toc(doc.add_paragraph())
    doc.add_page_break()

    doc.add_heading("Resumen de ejecución", level=1)
    _key_value_table(
        doc,
        [
            ("Video", data.video_name),
            ("Capturas", str(len(data.evidence))),
            ("Segmentos de transcripción", str(len(data.segments))),
            ("Origen de la transcripción", data.transcript_source_label),
            ("Perfil de procesamiento", data.profile_label),
            ("Correcciones manuales aplicadas", str(len(data.corrections))),
            ("Versión de la aplicación", data.app_version),
        ],
    )
    if data.metadata is not None:
        meta = data.metadata
        doc.add_heading("Metadatos del video", level=2)
        _key_value_table(
            doc,
            [
                ("Duración", format_hms(meta.duration_seconds)),
                ("Resolución", f"{meta.width} × {meta.height}"),
                ("Cuadros por segundo", f"{meta.fps:.2f}"),
                ("Códec", meta.codec or "No detectado"),
            ],
        )
    if data.settings_summary:
        doc.add_heading("Configuración relevante", level=2)
        _key_value_table(doc, list(data.settings_summary.items()))
    if data.warnings:
        doc.add_heading("Advertencias", level=2)
        for warning in data.warnings:
            doc.add_paragraph(safe_text(warning), style="List Bullet")
    doc.add_page_break()

    doc.add_heading("Recorrido visual sincronizado", level=1)
    doc.add_paragraph(
        "Cada sección corresponde al intervalo en el que la captura permaneció vigente. Se muestra el "
        "texto visible (OCR) y la transcripción que se solapa con ese intervalo."
    )
    for index, item in enumerate(data.evidence, start=1):
        doc.add_heading(f"{item.capture.id} · {item.start_display} a {item.end_display}", level=2)
        image_path = data.project_root / item.capture.file_rel
        size = _image_size(image_path) if image_path.exists() else None
        if size is not None:
            paragraph = doc.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.add_run().add_picture(str(image_path), width=Inches(size[0]), height=Inches(size[1]))
        else:
            doc.add_paragraph(f"[Captura no disponible: {Path(item.capture.file_rel).name}]")
        table = doc.add_table(rows=1, cols=2)
        table.style = "Table Grid"
        left, right = table.rows[0].cells
        for cell in (left, right):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
        left.paragraphs[0].add_run("Texto visible (OCR)").bold = True
        left.add_paragraph(safe_text(item.ocr_text) or "[Sin texto detectado]")
        right.paragraphs[0].add_run("Transcripción correspondiente").bold = True
        for segment in item.segments:
            _segment_paragraph(right.add_paragraph(), segment, data.include_sensitive)
        if not item.segments:
            right.add_paragraph("[Sin diálogo durante este intervalo]")
        if index < len(data.evidence):
            doc.add_page_break()

    timeline = build_timeline(data)
    if timeline:
        doc.add_page_break()
        doc.add_heading("Línea de tiempo segundo a segundo", level=1)
        doc.add_paragraph(
            "Una fila por segundo: lo que se dice y la pantalla vigente. El texto visible (OCR) se muestra al "
            "aparecer cada pantalla; el detalle repetido por segundo está en el Excel."
        )
        table = doc.add_table(rows=1, cols=4)
        table.style = "Table Grid"
        for cell, label in zip(
            table.rows[0].cells, ("Tiempo", "Lo que se dice", "Pantalla", "Texto visible (OCR)"), strict=True
        ):
            cell.paragraphs[0].add_run(label).bold = True
        for row in timeline:
            cells = table.add_row().cells
            cells[0].text = format_hms(float(row.second))
            cells[1].text = safe_text(row.spoken)
            cells[2].text = row.capture_id
            cells[3].text = safe_text(row.ocr_text) if row.new_screen else ""

    if data.segments:
        doc.add_page_break()
        doc.add_heading("Transcripción completa", level=1)
        for segment in data.segments:
            _segment_paragraph(doc.add_paragraph(), segment, data.include_sensitive, color=True)

    destination.mkdir(parents=True, exist_ok=True)
    path = destination / WORD_REPORT_NAME
    temporary = destination / f".{WORD_REPORT_NAME}.tmp"
    doc.save(str(temporary))
    temporary.replace(path)
    return path
